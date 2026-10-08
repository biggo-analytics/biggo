"""Writes the workbooks that the tests and the documentation read.

They are kept in the repository as they are; run this, with openpyxl installed, to change them:

    python3 testdata/make_workbooks.py
"""
import datetime
import pathlib

import openpyxl

root = pathlib.Path(__file__).resolve().parents[1]


def test_book(path):
    book = openpyxl.Workbook()

    sales = book.active
    sales.title = "Sales"
    sales.append(["Sales report"])
    sales.append(["printed", datetime.date(2026, 2, 1)])
    sales.append([])
    sales.append(["id", "branch ", "units", "price", "sold on", "at", "ok", "note", "took"])
    sales.append([1, "north", 10, 2.5, datetime.date(2026, 1, 5), datetime.datetime(2026, 1, 5, 9, 30), True, "ok", 90])
    sales.append([2, "south", None, 7.25, datetime.date(2026, 1, 6), datetime.datetime(2026, 1, 6, 17, 45, 30), False, "NA", 12.5])
    sales.append([3, "east", " 5 ", "NA", datetime.date(2026, 2, 28), datetime.datetime(2026, 2, 28), True, None, None])
    sales.append([])
    sales.append([4, "007", 12, 0.1, "2026-03-01", "2026-03-01T08:00:00", True, 42, datetime.timedelta(minutes=2)])
    sales["D9"].value = "#N/A"
    sales["D9"].data_type = "e"

    plain = book.create_sheet("Plain")
    for row in [[1, "north", 10], [2, "south", 7], [3, "east", 5]]:
        plain.append(row)

    thai = book.create_sheet("สาขา")
    thai.append(["ชื่อ", "จำนวน"])
    thai.append(["เชียงใหม่", 12])
    thai.append(["ภูเก็ต", 7])

    block = book.create_sheet("Block")
    block["A1"] = "notes on the side"
    block["C3"], block["D3"] = "code", "qty"
    block["C4"], block["D4"] = "W-01", 3
    block["C5"], block["D5"] = "G-07", 4
    block["C7"] = "total"
    block["D7"] = 7

    book.save(path)


def docs_book(path):
    book = openpyxl.Workbook()
    sales = book.active
    sales.title = "Sales"
    sales.append(["Sales by branch"])
    sales.append([])
    sales.append(["branch", "units", "price", "sold on", "paid"])
    sales.append(["north", 10, 2.5, datetime.date(2026, 1, 5), True])
    sales.append(["south", None, 7.25, datetime.date(2026, 1, 6), False])
    sales.append(["east", 5, "n/a", datetime.date(2026, 2, 28), True])

    targets = book.create_sheet("Targets")
    targets.append(["branch", "target"])
    for row in [["north", 8], ["south", 6], ["east", 9]]:
        targets.append(row)
    book.save(path)


test_book(root / "testdata/run/data/book.xlsx")
docs_book(root / "docs/data/branches.xlsx")
